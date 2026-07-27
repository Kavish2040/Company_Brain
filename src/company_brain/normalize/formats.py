"""The five M1 normalizers: markdown, PDF, docx, email, Slack export.

Each is a pure function of its input bytes. Where a format carries its own
timestamps (email Date:, Slack ts), we use those — never the file's mtime, which
changes on checkout and would break byte-identical re-ingestion.
"""

from __future__ import annotations

import json
import re
from datetime import UTC, datetime
from email import message_from_bytes
from email import utils as email_utils
from email.message import Message
from threading import Lock
from typing import Any

from company_brain.normalize.base import (
    Normalized,
    NormalizeError,
    Normalizer,
    Registry,
    decode,
    tidy,
)


class MarkdownNormalizer:
    """Markdown and plain text. Nearly a pass-through, deliberately."""

    name = "markdown"
    version = "1.0.0"
    suffixes = (".md", ".markdown", ".txt")

    def normalize(self, raw: bytes) -> Normalized:
        text = tidy(decode(raw))
        title = ""
        lines = text.split("\n")
        for line in lines:
            if line.startswith("# "):
                title = line[2:].strip()
                break
        if not title:
            title = next((line.strip() for line in lines if line.strip()), "Untitled")
        return Normalized(body=text, title=title[:200])


class PdfNormalizer:
    """PDF text extraction via pypdf.

    Text extraction is only *conditionally* deterministic: it is stable for a
    given pypdf version, and can shift across releases. That is exactly why the
    version below travels into the frontmatter and pypdf is pinned in the
    `converters` dependency group (ARCHITECTURE §4).
    """

    name = "pdf"
    version = "1.0.0"
    suffixes = (".pdf",)

    def normalize(self, raw: bytes) -> Normalized:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise NormalizeError(
                "pypdf is required for .pdf; install the 'converters' group"
            ) from exc

        import io

        try:
            reader = PdfReader(io.BytesIO(raw))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise NormalizeError(f"could not read pdf: {exc}") from exc

        body = tidy("\n\n".join(pages))
        meta: dict[str, object] = dict(reader.metadata or {})
        title = str(meta.get("/Title") or "").strip()
        if not title:
            first = next((ln.strip() for ln in body.split("\n") if ln.strip()), "")
            title = first[:200] or "Untitled PDF"
        author = str(meta.get("/Author") or "").strip()
        return Normalized(body=body, title=title, author_hints=(author,) if author else ())


# python-docx builds on lxml, whose element objects are not thread-safe. Under
# concurrent ingest this surfaced intermittently as
# `'lxml.etree._Element' object has no attribute 'Relationship_lst'` — a torn
# read of a partially-initialised part, not a corrupt file. One lock is cheaper
# than serialising the whole normalize phase, and docx is a small slice of any
# corpus.
_DOCX_LOCK = Lock()


class DocxNormalizer:
    """Word documents. Headings map to markdown levels; tables become pipe tables."""

    name = "docx"
    version = "1.0.0"
    suffixes = (".docx",)

    def normalize(self, raw: bytes) -> Normalized:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover
            raise NormalizeError(
                "python-docx is required for .docx; install the 'converters' group"
            ) from exc

        import io

        try:
            with _DOCX_LOCK:
                document = docx.Document(io.BytesIO(raw))
        except Exception as exc:
            raise NormalizeError(f"could not read docx: {exc}") from exc

        out: list[str] = []
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style else ""
            if style.startswith("heading"):
                digits = "".join(c for c in style if c.isdigit())
                level = min(int(digits), 6) if digits else 1
                out.append(f"{'#' * level} {text}")
            elif style.startswith("list"):
                out.append(f"- {text}")
            else:
                out.append(text)

        for table in document.tables:
            rows = [
                [cell.text.strip().replace("|", r"\|") for cell in r.cells] for r in table.rows
            ]
            if not rows:
                continue
            width = len(rows[0])
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "|".join([" --- "] * width) + "|")
            for row in rows[1:]:
                out.append("| " + " | ".join(row) + " |")

        body = tidy("\n\n".join(out))
        core = document.core_properties
        title = (core.title or "").strip()
        if not title:
            title = next(
                (ln.lstrip("# ").strip() for ln in body.split("\n") if ln.strip()),
                "Untitled Document",
            )
        author = (core.author or "").strip()
        return Normalized(
            body=body,
            title=title[:200],
            created=_as_utc(core.created),
            modified=_as_utc(core.modified),
            author_hints=(author,) if author else (),
        )


class EmailNormalizer:
    """RFC-822 messages. Prefers text/plain; falls back to a stripped text/html."""

    name = "email"
    version = "1.0.0"
    suffixes = (".eml",)

    def normalize(self, raw: bytes) -> Normalized:
        message = message_from_bytes(raw)
        subject = _header(message, "Subject") or "(no subject)"
        sender = _header(message, "From")
        recipients = [
            addr
            for header in ("To", "Cc")
            for _, addr in email_utils.getaddresses([_header(message, header) or ""])
            if addr
        ]

        sent: datetime | None = None
        date_header = _header(message, "Date")
        if date_header:
            try:
                parsed = email_utils.parsedate_to_datetime(date_header)
                sent = parsed.astimezone(UTC) if parsed.tzinfo else parsed.replace(tzinfo=UTC)
            except (TypeError, ValueError):
                sent = None

        header_block = [f"**From:** {sender}" if sender else ""]
        if recipients:
            header_block.append("**To:** " + ", ".join(sorted(set(recipients))))
        if sent:
            header_block.append(f"**Date:** {sent.isoformat()}")

        body = tidy(
            "\n\n".join([ln for ln in header_block if ln]) + "\n\n" + _email_body(message)
        )
        _, sender_addr = email_utils.parseaddr(sender or "")
        hints = tuple(h for h in (sender_addr, *sorted(set(recipients))) if h)
        return Normalized(
            body=body, title=subject[:200], created=sent, modified=sent, author_hints=hints
        )


class SlackExportNormalizer:
    """Slack export JSON — an array of message objects, one channel-day per file.

    Thread replies are nested under their parent so the markdown reads the way
    the channel did, rather than as a flat timestamp-ordered log.
    """

    name = "slack_export"
    version = "1.0.0"
    suffixes = (".json",)

    def normalize(self, raw: bytes) -> Normalized:
        try:
            payload = json.loads(decode(raw))
        except json.JSONDecodeError as exc:
            raise NormalizeError(f"not valid JSON: {exc}") from exc
        if not isinstance(payload, list):
            raise NormalizeError("slack export must be a JSON array of messages")

        messages = [m for m in payload if isinstance(m, dict) and m.get("type") == "message"]
        # Sorting by ts makes output independent of the array's order on disk.
        messages.sort(key=lambda m: str(m.get("ts", "")))

        roots = [m for m in messages if not m.get("thread_ts") or m["thread_ts"] == m["ts"]]
        replies: dict[str, list[dict[str, Any]]] = {}
        for m in messages:
            parent = m.get("thread_ts")
            if parent and parent != m.get("ts"):
                replies.setdefault(str(parent), []).append(m)

        lines: list[str] = []
        for root in roots:
            lines.append(_slack_line(root))
            for reply in replies.get(str(root.get("ts")), []):
                lines.append("  " + _slack_line(reply).replace("\n", "\n  "))

        stamps = sorted(t for t in (_slack_ts(m) for m in messages) if t is not None)
        authors = tuple(sorted({str(m["user"]) for m in messages if m.get("user")}))
        first = stamps[0] if stamps else None
        title = f"Slack thread — {first.date().isoformat()}" if first else "Slack thread"
        return Normalized(
            body=tidy("\n\n".join(lines)),
            title=title,
            created=first,
            modified=stamps[-1] if stamps else None,
            author_hints=authors,
        )


def _slack_line(message: dict[str, Any]) -> str:
    user = message.get("user_profile", {}).get("real_name") or message.get("user") or "unknown"
    stamp = _slack_ts(message)
    when = stamp.strftime("%H:%M") if stamp else "--:--"
    text = _slack_text(str(message.get("text", "")))
    return f"**{user}** ({when}): {text}"


def _slack_text(text: str) -> str:
    """Unwrap Slack's link and mention markup into readable text."""
    text = re.sub(r"<https?://[^|>]+\|([^>]+)>", r"\1", text)
    text = re.sub(r"<(https?://[^>]+)>", r"\1", text)
    text = re.sub(r"<@([A-Z0-9]+)\|([^>]+)>", r"@\2", text)
    text = re.sub(r"<@([A-Z0-9]+)>", r"@\1", text)
    text = re.sub(r"<#[A-Z0-9]+\|([^>]+)>", r"#\1", text)
    return text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")


def _slack_ts(message: dict[str, Any]) -> datetime | None:
    raw = message.get("ts")
    if raw is None:
        return None
    try:
        return datetime.fromtimestamp(float(raw), tz=UTC)
    except (TypeError, ValueError):
        return None


def _header(message: Message, name: str) -> str:
    value = message.get(name)
    if value is None:
        return ""
    from email.header import decode_header, make_header

    try:
        return str(make_header(decode_header(str(value))))
    except (UnicodeDecodeError, LookupError, ValueError):
        return str(value)


def _email_body(message: Message) -> str:
    plain, html = "", ""
    for part in message.walk() if message.is_multipart() else [message]:
        if part.get_content_maintype() == "multipart":
            continue
        if "attachment" in str(part.get("Content-Disposition", "")):
            continue
        payload = part.get_payload(decode=True)
        if not isinstance(payload, bytes):
            continue
        text = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
        if part.get_content_type() == "text/plain" and not plain:
            plain = text
        elif part.get_content_type() == "text/html" and not html:
            html = text
    if plain:
        return plain
    return re.sub(r"<[^>]+>", "", html) if html else ""


def _as_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value.astimezone(UTC) if value.tzinfo else value.replace(tzinfo=UTC)


def default_registry() -> Registry:
    registry = Registry()
    normalizers: tuple[Normalizer, ...] = (
        MarkdownNormalizer(),
        PdfNormalizer(),
        DocxNormalizer(),
        EmailNormalizer(),
        SlackExportNormalizer(),
    )
    for normalizer in normalizers:
        registry.register(normalizer)
    return registry
